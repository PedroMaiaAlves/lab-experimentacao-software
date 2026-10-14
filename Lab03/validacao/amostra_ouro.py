"""S02-A1: sorteio reprodutível da amostra-ouro e geração do modelo de rotulagem.

Uso (na raiz do projeto):
    python -m validacao.amostra_ouro
"""
import argparse
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Valores canônicos dos rótulos (a concordância da S02-A3 usa as mesmas listas)
TIPOS = ["biblioteca/framework", "aplicação/serviço", "ferramenta CLI", "outro"]
ENTREGAS = ["sim", "não", "incerto"]
CORRETIVA = ["sim", "não"]

COLS_REPOS = ["full_name", "repo_url", "tipo_projeto", "releases_sao_entregas"]
COLS_RELEASES = ["full_name", "release_id", "tag_name", "release_url", "compare_url", "corretiva"]
COLS_LEITURA_RELEASES = ["full_name", "id", "tag_name", "draft", "prerelease", "published_at", "html_url"]

INSTRUCOES = [
    "ROTULAGEM INDEPENDENTE: preencha sozinho, sem consultar os outros avaliadores e sem ver a saída da heurística.",
    "",
    "Aba 'repositorios' (uma linha por repositório):",
    "  tipo_projeto: biblioteca/framework (código usado como dependência por outros programas); "
    "aplicação/serviço (programa ou serviço usado diretamente); ferramenta CLI (programa usado pela linha de comando); "
    "outro (documentação, listas, tutoriais, dados, configurações, GitHub Actions).",
    "  releases_sao_entregas: sim se as releases entregam versões a usuários (binários, pacotes, notas de versão); "
    "não se são marcações automáticas ou sem entrega real; incerto se não der para decidir.",
    "",
    "Aba 'releases' (até 5 releases sorteadas por repositório):",
    "  corretiva: sim se o objetivo principal da release é corrigir defeito recente (correções, hotfix, revert); "
    "não nos demais casos. Use release_url (notas) e compare_url (lista de commits).",
    "",
    "Use as listas suspensas, não deixe células em branco e não altere as colunas de identificação.",
]

AZUL = PatternFill("solid", fgColor="DDEBF7")
AMARELO = PatternFill("solid", fgColor="FFF2CC")


def carregar_config(caminho):
    with open(caminho, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    return {"semente": int(cfg["semente"]),
            "inicio": str(cfg["janela_inicio"]),
            "fim": str(cfg["janela_fim"])}


def sortear_repositorios(amostra, n, semente):
    """n repositórios com semente fixa. Ordena por nome antes, para a ordem do CSV não mudar o sorteio."""
    ordenada = amostra.sort_values("full_name").reset_index(drop=True)
    if len(ordenada) < n:
        raise ValueError(f"A amostra tem {len(ordenada)} repositórios; são necessários {n}.")
    return ordenada.sample(n=n, random_state=semente).sort_values("full_name").reset_index(drop=True)


def releases_principais_na_janela(releases, inicio, fim):
    """Releases da definição principal: draft=false, prerelease=false e publicadas na janela (extremos incluídos)."""
    nao_draft = releases["draft"].astype(str).str.lower().eq("false")
    nao_pre = releases["prerelease"].astype(str).str.lower().eq("false")
    data = releases["published_at"].fillna("").astype(str).str[:10]
    return releases[nao_draft & nao_pre & (data >= inicio) & (data <= fim)]


def sortear_releases(principais, repos, n, semente):
    """Até n releases por repositório, com semente fixa (ordenadas por id antes de sortear)."""
    partes = []
    for nome in repos:
        sub = principais[principais["full_name"] == nome].copy()
        if sub.empty:
            continue
        sub["_id"] = pd.to_numeric(sub["id"])
        sub = sub.sort_values("_id")
        partes.append(sub.sample(n=min(n, len(sub)), random_state=semente).drop(columns="_id"))
    if not partes:
        return principais.iloc[0:0].copy()
    return pd.concat(partes, ignore_index=True)


def link_comparacao(repo, anterior, atual):
    if not anterior:
        return ""
    return f"https://github.com/{repo}/compare/{quote(anterior, safe='/')}...{quote(atual, safe='/')}"


def adicionar_links(sorteadas, comparacoes):
    df = sorteadas.rename(columns={"id": "release_id", "html_url": "release_url"}).copy()
    if comparacoes is not None and not comparacoes.empty:
        comp = comparacoes.drop_duplicates("release_id")[["release_id", "previous_tag_name"]]
        df = df.merge(comp, on="release_id", how="left")
        df["previous_tag_name"] = df["previous_tag_name"].fillna("")
    else:
        df["previous_tag_name"] = ""
    df["compare_url"] = [link_comparacao(r, p, t)
                         for r, p, t in zip(df["full_name"], df["previous_tag_name"], df["tag_name"])]
    cols = ["full_name", "release_id", "tag_name", "published_at", "release_url", "compare_url"]
    return df[cols].sort_values(["full_name", "published_at"]).reset_index(drop=True)


def _aba(wb, titulo, colunas, linhas, colunas_link, listas):
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = AZUL
    for linha in linhas:
        ws.append(linha)
    n = len(linhas)
    for nome in colunas_link:
        j = colunas.index(nome) + 1
        ws.column_dimensions[get_column_letter(j)].width = 45
        for i in range(2, n + 2):
            cel = ws.cell(row=i, column=j)
            if cel.value:
                cel.hyperlink = cel.value
                cel.font = Font(color="0563C1", underline="single")
    for nome, opcoes in listas.items():
        j = colunas.index(nome) + 1
        letra = get_column_letter(j)
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Valor inválido",
                            error="Escolha um valor da lista.")
        ws.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{n + 1}")
        ws.column_dimensions[letra].width = 24
        for i in range(2, n + 2):
            ws.cell(row=i, column=j).fill = AMARELO
    ws.column_dimensions["A"].width = 38
    ws.freeze_panes = "B2"


def escrever_modelo(repos, releases, caminho):
    """Modelo de rotulagem (xlsx). xlsx evita o problema do CSV com ';' no Excel em português."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "instrucoes"
    for i, linha in enumerate(INSTRUCOES, 1):
        ws0.cell(row=i, column=1, value=linha).alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 130
    linhas_repos = [[r.full_name, r.repo_url, None, None] for r in repos.itertuples()]
    _aba(wb, "repositorios", COLS_REPOS, linhas_repos, ["repo_url"],
         {"tipo_projeto": TIPOS, "releases_sao_entregas": ENTREGAS})
    linhas_rel = [[r.full_name, int(r.release_id), r.tag_name, r.release_url, r.compare_url, None]
                  for r in releases.itertuples()]
    _aba(wb, "releases", COLS_RELEASES, linhas_rel, ["release_url", "compare_url"],
         {"corretiva": CORRETIVA})
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Sorteia a amostra-ouro e gera o modelo de rotulagem.")
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--amostra", default="data/metadados.csv",
                    help="CSV da amostra final congelada (precisa da coluna full_name)")
    ap.add_argument("--releases", default="data/releases.csv")
    ap.add_argument("--comparacoes", default="data/comparacoes_releases.csv")
    ap.add_argument("--saida", default="data/amostra_ouro")
    ap.add_argument("--n-repos", type=int, default=60)
    ap.add_argument("--n-releases", type=int, default=5)
    a = ap.parse_args(argv)

    cfg = carregar_config(a.config)
    saida = Path(a.saida)
    saida.mkdir(parents=True, exist_ok=True)

    amostra = pd.read_csv(a.amostra, dtype=str, keep_default_na=False)
    repos = sortear_repositorios(amostra, a.n_repos, cfg["semente"])
    repos["repo_url"] = "https://github.com/" + repos["full_name"]

    rel = pd.read_csv(a.releases, usecols=COLS_LEITURA_RELEASES, dtype=str, keep_default_na=False)
    rel = rel[rel["full_name"].isin(repos["full_name"])]
    principais = releases_principais_na_janela(rel, cfg["inicio"], cfg["fim"])
    sorteadas = sortear_releases(principais, list(repos["full_name"]), a.n_releases, cfg["semente"])

    comparacoes = None
    if Path(a.comparacoes).exists():
        comparacoes = pd.read_csv(a.comparacoes, usecols=["release_id", "previous_tag_name"],
                                  dtype=str, keep_default_na=False)
    else:
        print(f"AVISO: {a.comparacoes} não encontrado; compare_url ficará vazio.")
    releases = adicionar_links(sorteadas, comparacoes)

    repos.to_csv(saida / "repos_sorteados.csv", index=False, encoding="utf-8")
    releases.to_csv(saida / "releases_sorteadas.csv", index=False, encoding="utf-8")
    escrever_modelo(repos, releases, saida / "modelo" / "rotulos_modelo.xlsx")

    contagem = releases.groupby("full_name").size().reindex(repos["full_name"]).fillna(0).astype(int)
    poucas = contagem[contagem < a.n_releases]
    print(f"{len(repos)} repositórios sorteados (semente {cfg['semente']}); {len(releases)} releases sorteadas.")
    if len(poucas):
        print(f"AVISO: {len(poucas)} repositório(s) com menos de {a.n_releases} releases principais na janela:")
        for nome, n in poucas.items():
            print(f"  {nome}: {n}")
    print(f"Arquivos em {saida}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())