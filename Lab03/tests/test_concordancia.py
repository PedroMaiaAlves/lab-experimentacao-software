import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from validacao.amostra_ouro import escrever_modelo
from validacao.concordancia import (
    ErroRotulos, calcular_kappa, consenso, executar, interpretar_kappa, ler_decisoes, main, tabela_dimensao)


def tab(linhas):
    return pd.DataFrame(linhas, columns=["A", "B", "C"], index=[f"r{i}" for i in range(len(linhas))])


def test_kappa_concordancia_perfeita_e_1():
    assert calcular_kappa(tab([["sim"] * 3, ["sim"] * 3, ["não"] * 3, ["não"] * 3])) == pytest.approx(1.0)


def test_kappa_calculado_a_mao_vale_menos_um_terco():
    # P_i = 1/3 em todos os itens; p_sim = p_não = 0,5; Pe = 0,5; kappa = (1/3 - 1/2) / (1/2)
    t = tab([["sim", "sim", "não"], ["sim", "sim", "não"], ["não", "não", "sim"], ["não", "não", "sim"]])
    assert calcular_kappa(t) == pytest.approx(-1 / 3)


def test_kappa_sem_variacao_e_indefinido():
    k = calcular_kappa(tab([["sim"] * 3] * 4))
    assert math.isnan(k) and interpretar_kappa(k).startswith("indefinido")


@pytest.mark.parametrize("k,esperado", [(-0.1, "pobre"), (0.0, "leve"), (0.2, "leve"), (0.35, "razoável"),
                                         (0.5, "moderada"), (0.7, "substancial"), (0.9, "quase perfeita")])
def test_interpretacao_landis_koch(k, esperado):
    assert interpretar_kappa(k) == esperado


def test_consenso_maioria_unanime_e_reuniao():
    t = tab([["sim", "sim", "sim"], ["sim", "não", "sim"], ["sim", "não", "incerto"]])
    c = consenso(t, "releases_sao_entregas", {})
    assert c["decisao"].tolist() == ["unanime", "maioria", "reuniao"]
    assert c["rotulo"].tolist() == ["sim", "sim", ""]
    c2 = consenso(t, "releases_sao_entregas", {("releases_sao_entregas", "r2"): "incerto"})
    assert c2.loc["r2", "rotulo"] == "incerto"


def test_adjudicacao_por_ia_nao_e_descrita_como_reuniao():
    t = tab([["sim", "não", "incerto"]])
    c = consenso(t, "releases_sao_entregas", {("releases_sao_entregas", "r0"): "incerto"}, "adjudicacao_IA")
    assert c.loc["r0", "decisao"] == "adjudicacao_IA"


def test_chaves_duplicadas_nao_sao_silenciosamente_removidas():
    d = {av: {"repositorios": pd.DataFrame({"full_name": ["o/r", "o/r"], "tipo_projeto": ["outro", "outro"]})}
         for av in "ABC"}
    with pytest.raises(ErroRotulos, match="duplicadas"):
        tabela_dimensao(d, "tipo_projeto")


def test_decisoes_com_rotulo_invalido_levantam_erro(tmp_path):
    (tmp_path / "d.csv").write_text("dimensao,chave,rotulo\ncorretiva,a|1,talvez\n")
    with pytest.raises(ErroRotulos):
        ler_decisoes(tmp_path / "d.csv")


def _preencher(origem, destino, tipos, entregas, corretivas):
    wb = load_workbook(origem)
    for i, (t, e) in enumerate(zip(tipos, entregas), start=2):
        wb["repositorios"].cell(row=i, column=3, value=t)
        wb["repositorios"].cell(row=i, column=4, value=e)
    for i, c in enumerate(corretivas, start=2):
        wb["releases"].cell(row=i, column=6, value=c)
    wb.save(destino)


@pytest.fixture
def pasta(tmp_path):
    repos = pd.DataFrame({"full_name": ["a/x", "b/y", "c/z"],
                          "repo_url": ["https://github.com/a/x", "https://github.com/b/y", "https://github.com/c/z"]})
    rels = pd.DataFrame({"full_name": ["a/x", "a/x", "b/y"], "release_id": ["11", "12", "21"],
                         "tag_name": ["v1", "v2", "v1"], "release_url": ["u1", "u2", "u3"],
                         "compare_url": ["", "c2", "c3"]})
    modelo = tmp_path / "modelo.xlsx"
    escrever_modelo(repos, rels, modelo)
    (tmp_path / "rotulos").mkdir()
    return tmp_path, modelo


def test_executar_de_ponta_a_ponta(pasta):
    p, modelo = pasta
    lib, app, cli = "biblioteca/framework", "aplicação/serviço", "ferramenta CLI"
    _preencher(modelo, p / "rotulos" / "rotulos_A.xlsx", [lib, app, cli], ["sim", "sim", "não"], ["sim", "não", "não"])
    _preencher(modelo, p / "rotulos" / "rotulos_B.xlsx", [lib, app, "outro"], ["sim", "não", "não"], ["sim", "não", "sim"])
    _preencher(modelo, p / "rotulos" / "rotulos_C.xlsx", [lib, cli, "outro"], ["sim", "incerto", "não"], ["sim", "não", "sim"])
    # c/z: A=CLI, B=outro, C=outro -> maioria; b/y: app, app, CLI -> maioria; sem 1-1-1 no tipo
    # releases_sao_entregas de b/y: sim, não, incerto -> 1-1-1 (precisa de decisão)
    assert executar(p, tex=p / "kappa.tex") == 1
    (p / "decisoes_reuniao.csv").write_text("dimensao,chave,rotulo\nreleases_sao_entregas,b/y,incerto\n")
    assert executar(p, tex=p / "kappa.tex") == 0

    repos = pd.read_csv(p / "consenso_repos.csv").set_index("full_name")
    assert repos.loc["c/z", "tipo_projeto"] == "outro" and repos.loc["c/z", "tipo_projeto_decisao"] == "maioria"
    assert repos.loc["b/y", "releases_sao_entregas"] == "incerto"
    rel = pd.read_csv(p / "consenso_releases.csv")
    assert list(rel["tag_name"]) == ["v1", "v2", "v1"] and rel["corretiva"].tolist() == ["sim", "não", "sim"]
    kappa = pd.read_csv(p / "kappa_fleiss.csv")
    assert set(kappa["dimensao"]) == {"tipo_projeto", "releases_sao_entregas", "corretiva"}
    assert "\\kappa" in (p / "kappa.tex").read_text(encoding="utf-8")
    assert (p / "casos_divergentes.csv").exists()


def test_celula_em_branco_e_valor_invalido_sao_reportados(pasta, capsys):
    p, modelo = pasta
    lib = "biblioteca/framework"
    _preencher(modelo, p / "rotulos" / "rotulos_A.xlsx", [lib, lib, lib], ["sim"] * 3, ["sim", "não", "não"])
    _preencher(modelo, p / "rotulos" / "rotulos_B.xlsx", [lib, None, lib], ["sim"] * 3, ["sim", "não", "não"])
    _preencher(modelo, p / "rotulos" / "rotulos_C.xlsx", [lib, lib, "livro"], ["sim"] * 3, ["sim", "não", "não"])
    assert main(["--pasta", str(p), "--tex", str(p / "k.tex")]) == 2
    saida = capsys.readouterr().out
    assert "B / tipo_projeto / linha 3" in saida and "'livro'" in saida


def test_arquivo_de_avaliador_ausente(pasta, capsys):
    p, _ = pasta
    assert main(["--pasta", str(p)]) == 2
    assert "não encontrado" in capsys.readouterr().out
