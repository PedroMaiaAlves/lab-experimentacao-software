import pandas as pd
import pytest
from openpyxl import load_workbook

from validacao.amostra_ouro import (
    adicionar_links, link_comparacao, main, releases_principais_na_janela, sortear_releases,
    sortear_repositorios)


def amostra(n):
    return pd.DataFrame({"full_name": [f"o/r{i:03d}" for i in range(n)], "stars": ["1500"] * n})


def releases_csv(caminho, repos, por_repo=8):
    linhas = []
    rid = 1000
    for r in repos:
        for k in range(por_repo):
            rid += 1
            linhas.append({
                "full_name": r, "id": rid, "tag_name": f"v1.{k}", "target_commitish": "main",
                "draft": "False", "prerelease": "True" if k == 0 else "False",
                "published_at": f"2026-0{k + 1}-10T10:00:00Z", "created_at": "x",
                "html_url": f"https://github.com/{r}/releases/tag/v1.{k}", "name": "n",
                "body": "linha 1\nlinha 2, com \"aspas\""})
    pd.DataFrame(linhas).to_csv(caminho, index=False)


def comparacoes_csv(caminho, repos, por_repo=8):
    linhas, rid = [], 1000
    for r in repos:
        for k in range(por_repo):
            rid += 1
            linhas.append({"full_name": r, "release_id": rid, "tag_name": f"v1.{k}",
                           "published_at": "x", "previous_tag_name": f"v1.{k - 1}" if k else "",
                           "status": "ok", "commits_count": 3})
    pd.DataFrame(linhas).to_csv(caminho, index=False)


def test_sorteio_de_repositorios_nao_depende_da_ordem_do_csv():
    base = amostra(100)
    embaralhada = base.sample(frac=1, random_state=1)
    a = sortear_repositorios(base, 60, 42)["full_name"].tolist()
    b = sortear_repositorios(embaralhada, 60, 42)["full_name"].tolist()
    assert a == b and len(a) == 60
    assert a != sortear_repositorios(base, 60, 7)["full_name"].tolist()


def test_amostra_pequena_demais_levanta_erro():
    with pytest.raises(ValueError):
        sortear_repositorios(amostra(10), 60, 42)


def test_releases_principais_excluem_draft_prerelease_e_fora_da_janela():
    df = pd.DataFrame({
        "draft": ["False", "True", "False", "False", "False"],
        "prerelease": ["False", "False", "True", "False", "False"],
        "published_at": ["2026-01-10T00:00:00Z", "2026-01-10T00:00:00Z", "2026-01-10T00:00:00Z",
                         "2025-09-30T23:59:59Z", "2026-09-30T23:59:59Z"]})
    ok = releases_principais_na_janela(df, "2025-10-01", "2026-09-30")
    assert list(ok.index) == [0, 4]


def test_sortear_releases_limita_a_n_e_e_reprodutivel():
    df = pd.DataFrame({"full_name": ["a/x"] * 8 + ["b/y"] * 3, "id": [str(i) for i in range(11)]})
    a = sortear_releases(df, ["a/x", "b/y"], 5, 42)
    b = sortear_releases(df.sample(frac=1, random_state=3), ["a/x", "b/y"], 5, 42)
    assert a.groupby("full_name").size().to_dict() == {"a/x": 5, "b/y": 3}
    assert sorted(a["id"]) == sorted(b["id"])


def test_link_de_comparacao():
    assert link_comparacao("o/r", "", "v1") == ""
    assert link_comparacao("o/r", "v1.0", "v1.1") == "https://github.com/o/r/compare/v1.0...v1.1"


def test_adicionar_links_sem_comparacoes_deixa_compare_vazio():
    s = pd.DataFrame({"full_name": ["a/x"], "id": ["1"], "tag_name": ["v1"],
                      "published_at": ["2026-01-01T00:00:00Z"], "html_url": ["https://u"]})
    out = adicionar_links(s, None)
    assert out.loc[0, "compare_url"] == "" and out.loc[0, "release_url"] == "https://u"


def test_main_gera_arquivos_e_modelo_xlsx(tmp_path, capsys):
    am = amostra(70)
    am.to_csv(tmp_path / "metadados.csv", index=False)
    releases_csv(tmp_path / "releases.csv", am["full_name"])
    comparacoes_csv(tmp_path / "comp.csv", am["full_name"])
    (tmp_path / "config.yaml").write_text("janela_inicio: 2025-10-01\njanela_fim: 2026-09-30\nsemente: 42\n")
    args = ["--config", str(tmp_path / "config.yaml"), "--amostra", str(tmp_path / "metadados.csv"),
            "--releases", str(tmp_path / "releases.csv"), "--comparacoes", str(tmp_path / "comp.csv"),
            "--saida", str(tmp_path / "ouro")]
    assert main(args) == 0
    repos = pd.read_csv(tmp_path / "ouro" / "repos_sorteados.csv")
    rels = pd.read_csv(tmp_path / "ouro" / "releases_sorteadas.csv")
    assert len(repos) == 60
    assert rels.groupby("full_name").size().max() == 5 and len(rels) == 300
    assert rels["compare_url"].str.startswith("https://github.com/").sum() >= 1

    wb = load_workbook(tmp_path / "ouro" / "modelo" / "rotulos_modelo.xlsx")
    assert wb.sheetnames == ["instrucoes", "repositorios", "releases"]
    assert wb["repositorios"].max_row == 61 and wb["releases"].max_row == 301
    assert len(wb["repositorios"].data_validations.dataValidation) == 2
    assert wb["repositorios"]["C2"].value is None  # rótulo em branco

    primeira = rels.iloc[0]
    assert primeira["release_url"].startswith("https://github.com/")
    # reprodutível: segunda execução gera exatamente os mesmos sorteios
    antes = rels.copy()
    assert main(args) == 0
    pd.testing.assert_frame_equal(antes, pd.read_csv(tmp_path / "ouro" / "releases_sorteadas.csv"))
    assert "60 repositórios sorteados" in capsys.readouterr().out